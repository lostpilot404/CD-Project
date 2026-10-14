# -*- coding: utf-8 -*-
"""Build the CodeTrace BCSE307P project report from the official VIT template.

Usage:
    python3 submission/tools/build_bcse307p_report.py [--pages toc_pages.json] [--out PATH]

The script opens ``BCSE307P_Complete_Project_Report_Template.docx`` (never
modified), fills the placeholders that describe the project, removes the
instructional text, inserts the chapter content defined in
``report_content.py`` and writes ``submission/CodeTrace_BCSE307P_Project_Report.docx``.

Pagination is not computed by python-docx, so the table of contents is
pre-rendered with page numbers taken from ``--pages`` (a JSON object mapping
heading text to page number, produced by ``render_check.mjs`` +
``measure_pages.py``).  Word is additionally asked to refresh all fields when
the document is opened (``w:updateFields``), which regenerates the TOC with
Word's own pagination.
"""
from __future__ import annotations

import argparse
import copy
import json
import os
import re
import sys
from typing import Dict, List, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, HERE)
import report_content as C  # noqa: E402

TEMPLATE = os.path.join(ROOT, "BCSE307P_Complete_Project_Report_Template.docx")
DEFAULT_OUT = os.path.join(ROOT, "submission", "CodeTrace_BCSE307P_Project_Report.docx")

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W14_NS = "http://schemas.microsoft.com/office/word/2010/wordml"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
PIC_NS = "http://schemas.openxmlformats.org/drawingml/2006/picture"

BODY_FONT_PT = 11
CODE_FONT = "Consolas"
TEXT_WIDTH_IN = 6.5
TOC_TAB_POS = 9350  # twips, right edge of the 6.5 in text column


# --------------------------------------------------------------------------- helpers
def w(tag: str) -> str:
    return "{%s}%s" % (W_NS, tag)


def strip_ids(el) -> None:
    """Remove w14:paraId / w14:textId from cloned paragraphs (must be unique)."""
    for e in el.iter():
        for attr in ("{%s}paraId" % W14_NS, "{%s}textId" % W14_NS):
            if attr in e.attrib:
                del e.attrib[attr]


def set_paragraph_text(p_el, text: str) -> None:
    """Keep the first run (and its formatting), set its text, drop other runs."""
    runs = p_el.findall(w("r"))
    if not runs:
        r = OxmlElement("w:r")
        p_el.append(r)
        runs = [r]
    first = runs[0]
    for extra in runs[1:]:
        p_el.remove(extra)
    for t in first.findall(w("t")):
        first.remove(t)
    for br in first.findall(w("br")):
        first.remove(br)
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    first.append(t)


def force_run_props(p_el, rPr_template) -> None:
    """Give every run of *p_el* the run properties of the template heading run."""
    for r in p_el.findall(w("r")):
        old = r.find(w("rPr"))
        if old is not None:
            r.remove(old)
        if rPr_template is not None:
            r.insert(0, copy.deepcopy(rPr_template))


def replace_in_runs(p_el, mapping: Dict[str, str]) -> int:
    count = 0
    for t in p_el.iter(w("t")):
        if t.text:
            new = t.text
            for old, rep in mapping.items():
                if old in new:
                    new = new.replace(old, rep)
                    count += 1
            if new != t.text:
                t.text = new
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return count


PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
             "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
             "sectPr", "pPrChange"]
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
               "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar",
               "tblLook", "tblCaption", "tblDescription", "tblPrChange"]


def reorder(el, order: List[str]) -> None:
    """Sort the children of *el* into the schema sequence given by *order*."""
    rank = {w(t): i for i, t in enumerate(order)}
    children = list(el)
    children.sort(key=lambda c: rank.get(c.tag, len(order)))
    for c in children:
        el.remove(c)
    for c in children:
        el.append(c)


class Inserter:
    """Insert new block elements after a moving anchor."""

    def __init__(self, anchor_el):
        self.cur = anchor_el

    def add(self, el):
        self.cur.addnext(el)
        self.cur = el
        return el


# --------------------------------------------------------------------------- inline markup
INLINE_RE = re.compile(r"(`[^`]*`|\*\*.+?\*\*|\*[^*]+?\*|\{[a-z]+\})")


class Citations:
    def __init__(self):
        self.order: List[str] = []

    def number(self, key: str) -> int:
        if key not in C.REFERENCES:
            raise KeyError("unknown citation key %r" % key)
        if key not in self.order:
            self.order.append(key)
        return self.order.index(key) + 1


CITES = Citations()


def add_runs(paragraph, text: str, size_pt: Optional[float] = None, bold=False, italic=False,
             code_size_pt: Optional[float] = None, color: Optional[RGBColor] = None) -> None:
    """Append runs for *text* with inline markup to *paragraph*."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            _plain(paragraph, text[pos:m.start()], size_pt, bold, italic, color)
        tok = m.group(0)
        if tok.startswith("`"):
            r = paragraph.add_run(tok[1:-1])
            r.font.name = CODE_FONT
            r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CODE_FONT)
            r.font.size = Pt(code_size_pt or ((size_pt or BODY_FONT_PT) - 1))
            r.bold = bold or None
            r.italic = italic or None
            if color is not None:
                r.font.color.rgb = color
        elif tok.startswith("**"):
            add_runs(paragraph, tok[2:-2], size_pt, True, italic, code_size_pt, color)
        elif tok.startswith("*"):
            add_runs(paragraph, tok[1:-1], size_pt, bold, True, code_size_pt, color)
        else:
            key = tok[1:-1]
            _plain(paragraph, "[%d]" % CITES.number(key), size_pt, bold, italic, color)
        pos = m.end()
    if pos < len(text):
        _plain(paragraph, text[pos:], size_pt, bold, italic, color)


def _plain(paragraph, text, size_pt, bold, italic, color):
    if not text:
        return
    r = paragraph.add_run(text)
    if size_pt:
        r.font.size = Pt(size_pt)
    if bold:
        r.bold = True
    if italic:
        r.italic = True
    if color is not None:
        r.font.color.rgb = color


# --------------------------------------------------------------------------- block builders
class Builder:
    def __init__(self, doc, h2_template_el, pagebreak_template_el):
        self.doc = doc
        self.h2_tpl = h2_template_el
        self.pb_tpl = pagebreak_template_el
        self.headings: List[tuple] = []  # (level, text, paragraph element)
        self.list_counter = 0
        self.numbering_part = doc.part.numbering_part
        self.decimal_abstract_id = self._ensure_decimal_abstract()

    # -- numbering -----------------------------------------------------------
    def _ensure_decimal_abstract(self) -> int:
        numbering = self.numbering_part.element
        ids = [int(a.get(w("abstractNumId"))) for a in numbering.findall(w("abstractNum"))]
        new_id = max(ids + [0]) + 1
        xml = (
            '<w:abstractNum xmlns:w="%s" w:abstractNumId="%d">'
            '<w:multiLevelType w:val="hybridMultilevel"/>'
            '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
            '<w:lvlText w:val="%%1."/><w:lvlJc w:val="left"/>'
            '<w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr></w:lvl>'
            '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/>'
            '<w:lvlText w:val="%%2."/><w:lvlJc w:val="left"/>'
            '<w:pPr><w:ind w:left="1440" w:hanging="360"/></w:pPr></w:lvl>'
            "</w:abstractNum>" % (W_NS, new_id)
        )
        el = etree.fromstring(xml)
        abstracts = numbering.findall(w("abstractNum"))
        if abstracts:
            abstracts[-1].addnext(el)
        else:
            numbering.insert(0, el)
        return new_id

    def new_num_id(self) -> int:
        numbering = self.numbering_part.element
        nums = numbering.findall(w("num"))
        ids = [int(n.get(w("numId"))) for n in nums]
        new_id = max(ids + [0]) + 1
        xml = (
            '<w:num xmlns:w="%s" w:numId="%d"><w:abstractNumId w:val="%d"/>'
            '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'
            % (W_NS, new_id, self.decimal_abstract_id)
        )
        el = etree.fromstring(xml)
        if nums:
            nums[-1].addnext(el)
        else:
            numbering.append(el)
        return new_id

    # -- paragraphs ----------------------------------------------------------
    def body_paragraph(self, text: str, align=WD_ALIGN_PARAGRAPH.JUSTIFY, keep_next=False,
                       space_after=8, line=1.5, size=None):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = align
        pf.line_spacing = line
        pf.space_after = Pt(space_after)
        pf.space_before = Pt(0)
        if keep_next:
            pf.keep_with_next = True
        add_runs(p, text, size_pt=size)
        return p._p

    def heading2(self, text: str):
        el = copy.deepcopy(self.h2_tpl)
        strip_ids(el)
        set_paragraph_text(el, text)
        return el

    def heading3(self, text: str):
        p = self.doc.add_paragraph(style="Heading 3")
        pf = p.paragraph_format
        pf.space_before = Pt(10)
        pf.space_after = Pt(4)
        pf.keep_with_next = True
        r = p.add_run(text)
        r.bold = True
        r.italic = False
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(0, 0, 0)
        return p._p

    def list_items(self, items: List[str], numbered: bool):
        num_id = self.new_num_id() if numbered else 1
        els = []
        for i, item in enumerate(items):
            p = self.doc.add_paragraph(style="List Paragraph")
            pPr = p._p.get_or_add_pPr()
            numPr = OxmlElement("w:numPr")
            ilvl = OxmlElement("w:ilvl")
            ilvl.set(qn("w:val"), "0")
            nid = OxmlElement("w:numId")
            nid.set(qn("w:val"), str(num_id))
            numPr.append(ilvl)
            numPr.append(nid)
            pPr.append(numPr)
            pf = p.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.line_spacing = 1.25
            pf.space_after = Pt(8 if i == len(items) - 1 else 3)
            pf.left_indent = Inches(0.5)
            pf.first_line_indent = Inches(-0.25)
            add_runs(p, item)
            reorder(pPr, PPR_ORDER)
            els.append(p._p)
        return els

    def caption(self, text: str, keep_next=False):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.space_before = Pt(4)
        pf.space_after = Pt(12)
        pf.line_spacing = 1.15
        if keep_next:
            pf.keep_with_next = True
            pf.space_after = Pt(4)
        label, _, rest = text.partition(": ")
        r = p.add_run(label + ":")
        r.bold = True
        r.font.size = Pt(10)
        add_runs(p, " " + rest, size_pt=10)
        return p._p

    def figure(self, path: str, caption: str, width_in: float):
        abspath = os.path.join(ROOT, path)
        if not os.path.exists(abspath):
            raise FileNotFoundError(abspath)
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.keep_with_next = True
        pf.space_before = Pt(6)
        pf.space_after = Pt(0)
        r = p.add_run()
        shape = r.add_picture(abspath, width=Inches(min(width_in, TEXT_WIDTH_IN)))
        # thin grey outline so that white screenshots are delimited on the page
        spPr = shape._inline.find(".//{%s}spPr" % PIC_NS)
        if spPr is not None:
            ln = etree.SubElement(spPr, "{%s}ln" % A_NS, w="6350")
            fill = etree.SubElement(ln, "{%s}solidFill" % A_NS)
            etree.SubElement(fill, "{%s}srgbClr" % A_NS, val="BFBFBF")
        els = [p._p, self.caption(caption)]
        return els

    def table(self, caption: str, header: List[str], rows: List[List[str]], widths: List[float]):
        cap = self.caption(caption, keep_next=True)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tbl = t._tbl
        tblPr = tbl.tblPr
        # width + fixed layout + borders + cell margins
        for tag in ("w:tblW", "w:tblLayout", "w:tblBorders", "w:tblCellMar"):
            for old in tblPr.findall(qn(tag)):
                tblPr.remove(old)
        total = int(sum(widths) * 1440)
        tblW = OxmlElement("w:tblW")
        tblW.set(qn("w:w"), str(total))
        tblW.set(qn("w:type"), "dxa")
        tblPr.append(tblW)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement("w:" + side)
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), "808080")
            borders.append(b)
        tblPr.append(borders)
        mar = OxmlElement("w:tblCellMar")
        for side, val in (("top", 30), ("left", 80), ("bottom", 30), ("right", 80)):
            m = OxmlElement("w:" + side)
            m.set(qn("w:w"), str(val))
            m.set(qn("w:type"), "dxa")
            mar.append(m)
        tblPr.append(mar)
        reorder(tblPr, TBLPR_ORDER)
        # grid
        grid = tbl.find(qn("w:tblGrid"))
        for gc, wd in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(wd * 1440)))

        def fill_cell(cell, text, bold=False, shade=None):
            cell.width = Inches(widths[cell_idx[0]])
            tcPr = cell._tc.get_or_add_tcPr()
            if shade:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), shade)
                tcPr.append(shd)
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.space_after = Pt(0)
            pf.space_before = Pt(0)
            pf.line_spacing = 1.0
            add_runs(p, text, size_pt=10, bold=bold, code_size_pt=9)

        cell_idx = [0]
        hdr = t.rows[0]
        trPr = hdr._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        trPr.append(th)
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
        for j, text in enumerate(header):
            cell_idx[0] = j
            fill_cell(hdr.cells[j], text, bold=True, shade="F2F2F2")
        for i, row in enumerate(rows):
            tr = t.rows[i + 1]
            trPr = tr._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            trPr.append(cs)
            for j, text in enumerate(row):
                cell_idx[0] = j
                fill_cell(tr.cells[j], text)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(6)
        spacer.paragraph_format.line_spacing = 1.0
        spacer_r = spacer.add_run("")
        spacer_r.font.size = Pt(6)
        return [cap, tbl, spacer._p]

    def code(self, lines: List[str]):
        els = []
        for i, line in enumerate(lines):
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.left_indent = Inches(0.3)
            pf.right_indent = Inches(0.3)
            pf.line_spacing = 1.0
            pf.space_after = Pt(8 if i == len(lines) - 1 else 0)
            pf.space_before = Pt(6 if i == 0 else 0)
            pf.keep_with_next = i < len(lines) - 1
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "F3F3F3")
            pPr.append(shd)
            reorder(pPr, PPR_ORDER)
            r = p.add_run(line if line else " ")
            r.font.name = CODE_FONT
            r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CODE_FONT)
            r.font.size = Pt(9.5)
            els.append(p._p)
        return els

    def pagebreak(self):
        el = copy.deepcopy(self.pb_tpl)
        strip_ids(el)
        return el

    # -- dispatch ------------------------------------------------------------
    def render_blocks(self, blocks, ins: Inserter, existing_h2: Dict[str, object]):
        for block in blocks:
            kind = block[0]
            if kind == "h2":
                text = block[1]
                num = text.split(" ", 1)[0]
                if num in existing_h2:
                    el = existing_h2[num]
                    set_paragraph_text(el, text)
                    ins.cur = el
                else:
                    el = ins.add(self.heading2(text))
                self.headings.append((2, text, el))
            elif kind == "h3":
                ins.add(self.heading3(block[1]))
            elif kind == "p":
                ins.add(self.body_paragraph(block[1]))
            elif kind == "bullets":
                for el in self.list_items(block[1], numbered=False):
                    ins.add(el)
            elif kind == "numbered":
                for el in self.list_items(block[1], numbered=True):
                    ins.add(el)
            elif kind == "figure":
                for el in self.figure(block[1], block[2], block[3]):
                    ins.add(el)
            elif kind == "table":
                for el in self.table(block[1], block[2], block[3], block[4]):
                    ins.add(el)
            elif kind == "code":
                for el in self.code(block[1]):
                    ins.add(el)
            elif kind == "caption":
                ins.add(self.caption(block[1]))
            elif kind == "pagebreak":
                ins.add(self.pagebreak())
            else:
                raise ValueError("unknown block %r" % (kind,))


# --------------------------------------------------------------------------- TOC / styles / settings
def ensure_toc_styles(doc) -> None:
    styles_el = doc.styles.element
    existing = {s.get(w("styleId")) for s in styles_el.findall(w("style"))}
    for sid, name, indent in (("TOC1", "toc 1", 0), ("TOC2", "toc 2", 220)):
        if sid in existing:
            continue
        xml = (
            '<w:style xmlns:w="%s" w:type="paragraph" w:styleId="%s">'
            '<w:name w:val="%s"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/>'
            '<w:autoRedefine/><w:uiPriority w:val="39"/><w:unhideWhenUsed/>'
            '<w:pPr><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="%d"/></w:tabs>'
            '<w:spacing w:after="100" w:line="276" w:lineRule="auto"/>%s</w:pPr>'
            "</w:style>" % (W_NS, sid, name, TOC_TAB_POS,
                            ('<w:ind w:left="%d"/>' % indent) if indent else "")
        )
        styles_el.append(etree.fromstring(xml))


def add_bookmark(p_el, bid: int, name: str) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    pPr = p_el.find(w("pPr"))
    if pPr is not None:
        pPr.addnext(start)
    else:
        p_el.insert(0, start)
    p_el.append(end)


def toc_entry_runs(text: str, anchor: str, page: str) -> List:
    def run(child):
        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        rPr.append(OxmlElement("w:noProof"))
        r.append(rPr)
        r.append(child)
        return r

    def t(txt):
        el = OxmlElement("w:t")
        el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        el.text = txt
        return el

    def fld(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el

    instr = OxmlElement("w:instrText")
    instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr.text = " PAGEREF %s \\h " % anchor
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("w:anchor"), anchor)
    hl.set(qn("w:history"), "1")
    for child in (t(text), OxmlElement("w:tab"), fld("begin"), instr, fld("separate"), t(page), fld("end")):
        hl.append(run(child))
    return hl


def build_toc(doc, headings, pages: Dict[str, str]) -> None:
    body = doc.element.body
    sdt = body.find(w("sdt"))
    content = sdt.find(w("sdtContent"))
    paras = content.findall(w("p"))
    first, last = paras[0], paras[-1]
    # remove any paragraphs between first and last (there are none in the template)
    for p in paras[1:-1]:
        content.remove(p)
    end_run = None
    for r in last.findall(w("r")):
        fc = r.find(w("fldChar"))
        if fc is not None and fc.get(qn("w:fldCharType")) == "end":
            end_run = r
    content.remove(last)
    cur = first
    for idx, (level, text, p_el) in enumerate(headings):
        anchor = "_Toc%06d" % (100001 + idx)
        add_bookmark(p_el, 500 + idx, anchor)
        if idx == 0:
            para = first
        else:
            para = OxmlElement("w:p")
            cur.addnext(para)
        pPr = para.find(w("pPr"))
        if pPr is None:
            pPr = OxmlElement("w:pPr")
            para.insert(0, pPr)
        for old in pPr.findall(w("pStyle")):
            pPr.remove(old)
        ps = OxmlElement("w:pStyle")
        ps.set(qn("w:val"), "TOC1" if level == 1 else "TOC2")
        pPr.insert(0, ps)
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:leader"), "dot")
        tab.set(qn("w:pos"), str(TOC_TAB_POS))
        tabs.append(tab)
        pPr.append(tabs)
        rPr = OxmlElement("w:rPr")
        rPr.append(OxmlElement("w:noProof"))
        pPr.append(rPr)
        page = str(pages.get(text, "00"))
        para.append(toc_entry_runs(text, anchor, page))
        cur = para
    if end_run is not None:
        cur.append(end_run)


def set_update_fields(doc) -> None:
    settings = doc.settings.element
    if settings.find(w("updateFields")) is not None:
        return
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    after = ["hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids",
             "mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping",
             "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade", "captions",
             "readModeInkLockDown", "smartTagType", "schemaLibrary", "shapeDefaults",
             "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"]
    for tag in after:
        el = settings.find(w(tag))
        if el is not None:
            el.addprevious(uf)
            return
    settings.append(uf)


# --------------------------------------------------------------------------- main
def expect(paragraphs, idx: int, contains: str) -> None:
    txt = paragraphs[idx].text
    if contains not in txt:
        raise AssertionError("template paragraph %d expected to contain %r, found %r" % (idx, contains, txt[:80]))


def build(out_path: str, pages: Dict[str, str]) -> None:
    doc = Document(TEMPLATE)
    ps = doc.paragraphs
    if len(ps) != 124:
        raise AssertionError("unexpected template layout: %d paragraphs" % len(ps))

    # --- sanity checks on the template map
    expect(ps, 8, "[PROJECT TITLE]")
    expect(ps, 23, "[PROJECT TITLE]")
    expect(ps, 30, "[Faculty Guide Name]")
    expect(ps, 37, "ABSTRACT")
    expect(ps, 38, "[150")
    expect(ps, 43, "CHAPTER 1")
    expect(ps, 45, "INTRODUCTION")
    expect(ps, 46, "1.1 Background")
    expect(ps, 52, "1.2 Problem Statement")
    expect(ps, 59, "1.3 Objectives")
    expect(ps, 66, "CHAPTER 2")
    expect(ps, 67, "LITERATURE SURVEY")
    expect(ps, 68, "2.1 Existing Systems")
    expect(ps, 76, "2.2 Research Gap")
    expect(ps, 79, "CHAPTER 3")
    expect(ps, 80, "SYSTEM DESIGN")
    expect(ps, 81, "3.1 Proposed System Architecture")
    expect(ps, 84, "3.2 Module Description")
    expect(ps, 88, "3.3 Compiler Phases Implementation")
    expect(ps, 102, "CHAPTER 4")
    expect(ps, 104, "Implementation Details")
    expect(ps, 120, "REFERENCES")
    expect(ps, 121, "[[1]")
    expect(ps, 122, "[[2]")
    if not ps[65]._p.findall(".//" + w("br")) or not ps[78]._p.findall(".//" + w("br")):
        raise AssertionError("page-break paragraphs not where expected")

    el = {i: ps[i]._p for i in range(len(ps))}

    # --- cover, bonafide, declaration
    set_paragraph_text(el[8], C.TITLE)
    rPr = el[8].find(w("r")).find(w("rPr"))
    if rPr is not None and rPr.find(w("b")) is None:
        rPr.insert(0, OxmlElement("w:bCs"))
        rPr.insert(0, OxmlElement("w:b"))
    replace_in_runs(el[23], {"[PROJECT TITLE]": "“" + C.TITLE + "”"})
    replace_in_runs(el[30], {"[PROJECT TITLE]": "“" + C.TITLE + "”",
                             "[Faculty Guide Name]": C.GUIDE})

    # --- chapter title case for chapter 4 (template uses capitals elsewhere)
    set_paragraph_text(el[104], "IMPLEMENTATION DETAILS")

    # --- normalise H2 headings that carry instruction text / trailing runs
    set_paragraph_text(el[81], "3.1 Proposed System Architecture")
    set_paragraph_text(el[84], "3.2 Module Description")
    set_paragraph_text(el[88], "3.3 Compiler Phases Implementation")
    set_paragraph_text(el[59], "1.3 Objectives & Novelty")

    builder = Builder(doc, h2_template_el=el[46], pagebreak_template_el=el[65])
    ensure_toc_styles(doc)
    h2_rPr = el[46].find(w("r")).find(w("rPr"))
    for i in (46, 52, 59, 68, 76, 81, 84, 88):
        force_run_props(el[i], h2_rPr)

    # --- abstract
    ins = Inserter(el[37])
    for para in C.ABSTRACT:
        ins.add(builder.body_paragraph(para))
    ins.add(builder.body_paragraph("**Keywords:** " + C.KEYWORDS))

    # --- chapters
    existing_h2 = {"1.1": el[46], "1.2": el[52], "1.3": el[59], "2.1": el[68], "2.2": el[76],
                   "3.1": el[81], "3.2": el[84], "3.3": el[88]}
    h1 = [(43, 45), (66, 67), (79, 80), (102, 104)]
    chapter_blocks = [C.CHAPTER1, C.CHAPTER2, C.CHAPTER3, C.CHAPTER4]
    for (num_idx, title_idx), blocks in zip(h1, chapter_blocks):
        builder.headings.append((1, ps[num_idx].text.strip(), el[num_idx]))
        builder.headings.append((1, ps[title_idx].text.strip() if title_idx != 104 else "IMPLEMENTATION DETAILS", el[title_idx]))
        ins = Inserter(el[title_idx])
        builder.render_blocks(blocks, ins, existing_h2)

    # page breaks before chapter 4 and references
    el[102].addprevious(builder.pagebreak())
    el[120].addprevious(builder.pagebreak())

    # --- references
    builder.headings.append((1, "REFERENCES", el[120]))
    ins = Inserter(el[120])
    missing = [k for k in C.REFERENCES if k not in CITES.order]
    if missing:
        raise AssertionError("references never cited in the text: %s" % missing)
    for n, key in enumerate(CITES.order, start=1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.left_indent = Inches(0.45)
        pf.first_line_indent = Inches(-0.45)
        pf.line_spacing = 1.15
        pf.space_after = Pt(6)
        pf.tab_stops.add_tab_stop(Inches(0.45))
        p.add_run("[%d]\t" % n)
        add_runs(p, C.REFERENCES[key])
        ins.add(p._p)

    # --- remove placeholders / instructions / empty heading paragraphs
    remove = [38, 44] + list(range(47, 52)) + list(range(53, 59)) + list(range(60, 65)) \
        + list(range(69, 76)) + [77, 82, 83, 85, 86, 87] + list(range(89, 102)) + [103] \
        + list(range(105, 120)) + [121, 122]
    for i in remove:
        p_el = el[i]
        p_el.getparent().remove(p_el)

    # --- TOC with pre-rendered entries, bookmarks on headings
    build_toc(doc, builder.headings, pages)
    set_update_fields(doc)

    # core properties
    doc.core_properties.title = C.TITLE
    doc.core_properties.subject = "BCSE307P – Compiler Design Lab Project Report"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.keywords = C.KEYWORDS

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    doc.save(out_path)
    print("wrote", out_path)
    print("headings:", len(builder.headings), "references:", len(CITES.order))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--pages", help="JSON mapping heading text -> page number")
    ap.add_argument("--out", default=DEFAULT_OUT)
    args = ap.parse_args()
    pages: Dict[str, str] = {}
    if args.pages:
        with open(args.pages, "r", encoding="utf-8") as fh:
            pages = {k: str(v) for k, v in json.load(fh).items()}
    build(args.out, pages)


if __name__ == "__main__":
    main()
